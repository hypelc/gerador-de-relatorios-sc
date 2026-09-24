"""Importadores e deteccao conservadora do contrato de entrada."""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
from pathlib import Path
from typing import Iterable, Literal

import openpyxl

from .models import (
    Diagnostic,
    ImportResult,
    IndicatorType,
    RecognizedTable,
    Series,
    SheetSummary,
    SourceSheet,
)


_YEAR_RE = re.compile(r"^(?:19|20)\d{2}$")
_MAP_CODES = {"4210", "4211", "4213", "4214", "4215", "4216", "4217", "4218"}


def _normalise(value: object) -> str:
    import unicodedata

    text = "" if value is None else str(value)
    return "".join(
        char
        for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )


def _as_year(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        number = int(value)
        return number if float(value) == number and 1900 <= number <= 2100 else None
    text = str(value).strip() if value is not None else ""
    return int(text) if _YEAR_RE.fullmatch(text) else None


def _coerce_csv_value(value: str) -> object:
    text = value.strip()
    if not text:
        return None
    if _YEAR_RE.fullmatch(text):
        return int(text)
    number_text = text.replace(",", ".") if "," in text and "." not in text else text
    try:
        number = float(number_text)
    except ValueError:
        return text
    return int(number) if number.is_integer() else number


def _cell_label(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _indicator_type(sheet: str, title: str) -> tuple[IndicatorType, str]:
    key = _normalise(f"{sheet} {title}").upper()
    if "MORTALIDADE" in key:
        return "mortalidade", "Mortalidade infantil (valor absoluto)"
    if "NASCIDOS" in key:
        return "nascidos", "Nascidos vivos (numero)"
    if "VACINA" in key or "IMUNIZ" in key or "COBERTURA" in key:
        return "vacina", "Cobertura vacinal (%)"
    # O contrato inicial e composto por tabelas de cobertura; mantemos essa
    # sugestao para nomes de aba curtos como BCG, PENTA e VARICELA. A usuaria
    # ainda pode trocar o tipo na etapa de mapeamento.
    return "vacina", "Cobertura vacinal (%)"


def _title_from_rows(sheet_name: str, rows: tuple[tuple[object, ...], ...], header_index: int) -> str:
    for row in reversed(rows[:header_index]):
        for value in row:
            text = _cell_label(value)
            if text and (" por Ano" in text or "POR ANO" in text.upper()):
                return text.split(" por Ano", 1)[0].strip()
    return sheet_name.strip()


def _year_cells(row: tuple[object, ...]) -> list[tuple[int, int]]:
    return [
        (column, year)
        for column, value in enumerate(row)
        if (year := _as_year(value)) is not None
    ]


def _find_label_column(rows: tuple[tuple[object, ...], ...], header_index: int, first_year_column: int) -> int:
    for column in range(first_year_column - 1, -1, -1):
        if any(_cell_label(row[column] if column < len(row) else None) for row in rows[header_index + 1 :]):
            return column
    return max(0, first_year_column - 1)


def _cell_reference(row: int, column: int) -> str:
    number = column + 1
    letters = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"{letters}{row + 1}"


def _is_blank_row(row: tuple[object, ...]) -> bool:
    return not any(value is not None and str(value).strip() for value in row)


def _is_formula_cell(source_sheet: SourceSheet, row: int, column: int) -> bool:
    return _cell_reference(row, column) in source_sheet.formulas


def _parse_table(
    source_sheet: SourceSheet,
    header_index: int,
    label_column: int,
    table_id: str,
    *,
    unit_override: str = "",
    indicator_override: IndicatorType | None = None,
    next_header_index: int | None = None,
) -> RecognizedTable:
    rows = source_sheet.rows
    header = rows[header_index] if header_index < len(rows) else ()
    year_cells = _year_cells(header)
    diagnostics: list[Diagnostic] = []
    if len(year_cells) < 3:
        diagnostics.append(
            Diagnostic(
                "error",
                "YEARS_TOO_FEW",
                "A tabela precisa de pelo menos tres anos validos.",
                f"Aba {source_sheet.name}, linha {header_index + 1}",
                "Escolha uma linha de cabecalho que contenha tres ou mais anos.",
            )
        )
    years = tuple(year for _, year in year_cells)
    columns = tuple(column for column, _ in year_cells)
    if label_column in columns:
        diagnostics.append(
            Diagnostic(
                "error",
                "LABEL_COLUMN_IS_YEAR",
                "A coluna de identificacao nao pode ser uma das colunas de ano.",
                f"Aba {source_sheet.name}, linha {header_index + 1}, coluna {label_column + 1}",
                "Escolha uma coluna de identificacao fora do conjunto de anos.",
            )
        )
    if len(set(years)) != len(years) or list(years) != sorted(years):
        diagnostics.append(
            Diagnostic(
                "error",
                "YEARS_NOT_ORDERED",
                "Os anos do cabecalho estao duplicados ou fora de ordem.",
                f"Aba {source_sheet.name}, linha {header_index + 1}",
                "Mantenha cada ano uma vez e em ordem crescente.",
            )
        )

    end = next_header_index if next_header_index is not None else len(rows)
    series: list[Series] = []
    seen_labels: set[str] = set()
    for row_index in range(header_index + 1, end):
        row = rows[row_index]
        label = _cell_label(row[label_column] if label_column < len(row) else None)
        values: list[float | None] = []
        nonempty = False
        invalid_cells: list[str] = []
        has_numeric = False
        for column in columns:
            value = row[column] if column < len(row) else None
            if value is not None and value != "":
                nonempty = True
            if value is None or value == "":
                values.append(None)
                if _is_formula_cell(source_sheet, row_index, column):
                    diagnostics.append(
                        Diagnostic(
                            "warning",
                            "FORMULA_NO_CACHED_VALUE",
                            "A formula nao possui valor calculado salvo no arquivo; o ponto foi mantido como ausente.",
                            f"Aba {source_sheet.name}, celula {_cell_reference(row_index, column)}",
                            "Abra e salve a planilha em um programa de planilhas antes de importar novamente.",
                        )
                    )
                continue
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
                invalid_cells.append(_cell_reference(row_index, column))
                values.append(None)
                continue
            has_numeric = True
            values.append(float(value))
        if not label or not nonempty:
            continue
        if invalid_cells:
            diagnostics.append(
                Diagnostic(
                    "error",
                    "INVALID_VALUE",
                    "Ha valores nao numericos na regiao de dados confirmada; eles nao foram convertidos silenciosamente.",
                    f"Aba {source_sheet.name}, linha {row_index + 1}, celulas {', '.join(invalid_cells)}",
                    "Corrija as celulas ou ajuste o cabecalho e a coluna de identificacao.",
                )
            )
        if not has_numeric and not invalid_cells:
            continue
        normalized_label = _normalise(label).upper()
        if normalized_label in seen_labels:
            diagnostics.append(
                Diagnostic(
                    "error",
                    "DUPLICATE_SERIES",
                    "A identificacao da serie aparece mais de uma vez na tabela.",
                    f"Aba {source_sheet.name}, linha {row_index + 1}: {label}",
                    "Remova a duplicata ou ajuste a selecao de tabela.",
                )
            )
        seen_labels.add(normalized_label)
        series.append(Series(label, tuple(values), row_index + 1))

    if not series:
        diagnostics.append(
            Diagnostic(
                "error",
                "TABLE_WITHOUT_SERIES",
                "Nenhuma serie numerica foi encontrada abaixo do cabecalho.",
                f"Aba {source_sheet.name}, linha {header_index + 1}",
                "Confira o cabecalho e a coluna que identifica as series.",
            )
        )

    title = _title_from_rows(source_sheet.name, rows, header_index)
    detected_type, detected_unit = _indicator_type(source_sheet.name, title)
    indicator_type = indicator_override or detected_type
    unit = unit_override.strip() or detected_unit
    code_count = sum(serie.code in _MAP_CODES for serie in series)
    confidence = min(1.0, 0.35 + (0.25 if len(years) >= 3 else 0) + (0.25 if series else 0) + (0.15 if code_count >= 8 else 0))
    return RecognizedTable(
        table_id=table_id,
        source_sheet=source_sheet.name,
        title=title,
        header_row=header_index + 1,
        label_column=label_column + 1,
        years=years,
        year_columns=tuple(column + 1 for column in columns),
        series=tuple(series),
        unit=unit,
        indicator_type=indicator_type,
        diagnostics=tuple(diagnostics),
        confidence=confidence,
    )


def _candidate_header_rows(source_sheet: SourceSheet) -> Iterable[int]:
    previous_candidate: int | None = None
    for index, row in enumerate(source_sheet.rows):
        if len(_year_cells(row)) >= 3:
            if previous_candidate is not None:
                separated = any(
                    _is_blank_row(source_sheet.rows[row_index])
                    for row_index in range(previous_candidate + 1, index)
                )
                if not separated:
                    continue
            previous_candidate = index
            yield index


def _parse_source(path: Path) -> tuple[tuple[SourceSheet, ...], Literal["xlsx", "csv"], int, list[Diagnostic]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        values_wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        formulas_wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
        source_sheets: list[SourceSheet] = []
        formula_count = 0
        for values_sheet, formulas_sheet in zip(values_wb.worksheets, formulas_wb.worksheets):
            rows = tuple(tuple(row) for row in values_sheet.iter_rows(values_only=True))
            formulas: dict[str, str] = {}
            for row in formulas_sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas[cell.coordinate] = cell.value
            formula_count += len(formulas)
            source_sheets.append(SourceSheet(values_sheet.title, rows, formulas))
        values_wb.close()
        formulas_wb.close()
        return tuple(source_sheets), "xlsx", formula_count, []
    if suffix == ".csv":
        raw = path.read_text(encoding="utf-8-sig")
        sample = raw[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t,")
        except csv.Error:
            dialect = csv.excel
        rows = tuple(tuple(_coerce_csv_value(value) for value in row) for row in csv.reader(io.StringIO(raw), dialect))
        return (SourceSheet(path.stem, rows),), "csv", 0, []
    raise ValueError("Formato nao suportado. Escolha um arquivo .xlsx ou .csv.")


def inspect_file(path: str | Path) -> ImportResult:
    source_path = Path(path).expanduser().resolve()
    if not source_path.exists() or not source_path.is_file():
        raise FileNotFoundError(f"Arquivo nao encontrado: {source_path}")
    if source_path.suffix.lower() not in {".xlsx", ".csv"}:
        raise ValueError("Formato nao suportado. Escolha um arquivo .xlsx ou .csv.")

    sheets, file_format, formula_count, diagnostics = _parse_source(source_path)
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    tables: list[RecognizedTable] = []
    summaries: list[SheetSummary] = []
    for sheet in sheets:
        header_rows = list(_candidate_header_rows(sheet))
        sheet_table_count = 0
        for index, header_index in enumerate(header_rows):
            header = sheet.rows[header_index]
            year_cells = _year_cells(header)
            if len(year_cells) < 3:
                continue
            first_year_column = year_cells[0][0]
            label_column = _find_label_column(sheet.rows, header_index, first_year_column)
            next_header = header_rows[index + 1] if index + 1 < len(header_rows) else None
            table_id = hashlib.sha1(
                f"{source_hash}:{sheet.name}:{header_index + 1}:{label_column + 1}".encode("utf-8")
            ).hexdigest()[:14]
            table = _parse_table(sheet, header_index, label_column, table_id, next_header_index=next_header)
            tables.append(table)
            sheet_table_count += 1
        if sheet_table_count:
            summaries.append(SheetSummary(sheet.name, sheet_table_count, "reconhecida", f"{sheet_table_count} tabela(s) candidata(s)."))
        else:
            diagnostic = Diagnostic(
                "warning",
                "SHEET_UNRECOGNIZED",
                "Aba auxiliar ou sem uma sequencia de pelo menos tres anos validos; nenhuma tabela sera usada automaticamente.",
                f"Aba {sheet.name}",
                "Confira se existe uma linha com anos e uma coluna de identificacao.",
            )
            diagnostics.append(diagnostic)
            summaries.append(SheetSummary(sheet.name, 0, "nao reconhecida", diagnostic.message))

    if not tables:
        diagnostics.append(
            Diagnostic(
                "error",
                "NO_VALID_TABLE",
                "Nenhuma tabela anual compativel foi encontrada no arquivo.",
                source_path.name,
                "Use .xlsx ou .csv com tres ou mais anos e valores numericos nas linhas seguintes.",
            )
        )
    return ImportResult(
        source_path=source_path,
        file_format=file_format,
        source_sha256=source_hash,
        tables=tuple(tables),
        sheets=sheets,
        sheet_summaries=tuple(summaries),
        diagnostics=tuple(diagnostics),
        formula_count=formula_count,
    )


def remap_table(
    result: ImportResult,
    table_id: str,
    *,
    header_row: int | None = None,
    label_column: int | None = None,
    unit: str = "",
    indicator_type: IndicatorType | None = None,
) -> ImportResult:
    original = result.table(table_id)
    sheet = next(sheet for sheet in result.sheets if sheet.name == original.source_sheet)
    selected_header = (header_row or original.header_row) - 1
    selected_label = (label_column or original.label_column) - 1
    if selected_header < 0 or selected_header >= len(sheet.rows):
        raise ValueError("Linha de cabecalho fora do intervalo da aba.")
    if selected_label < 0:
        raise ValueError("Coluna de identificacao invalida.")
    candidate_rows = set(_candidate_header_rows(sheet))
    if selected_header not in candidate_rows:
        raise ValueError("A linha escolhida nao foi reconhecida como cabecalho separado de outro bloco de dados.")
    next_header_index = next(
        (candidate for candidate in sorted(candidate_rows) if candidate > selected_header),
        None,
    )
    replacement = _parse_table(
        sheet,
        selected_header,
        selected_label,
        table_id,
        unit_override=unit,
        indicator_override=indicator_type,
        next_header_index=next_header_index,
    )
    tables = tuple(replacement if table.table_id == table_id else table for table in result.tables)
    return ImportResult(
        source_path=result.source_path,
        file_format=result.file_format,
        source_sha256=result.source_sha256,
        tables=tables,
        sheets=result.sheets,
        sheet_summaries=result.sheet_summaries,
        diagnostics=result.diagnostics,
        formula_count=result.formula_count,
    )
