from __future__ import annotations

import threading
from pathlib import Path

import openpyxl
import pytest
from pypdf import PdfReader

try:
    from gerador_sc.engine import OperationCancelled, ReportService
except ImportError:  # compatibilidade durante a primeira etapa red do teste
    from gerador_sc.engine import ReportService

    class OperationCancelled(RuntimeError):
        pass
from gerador_sc.models import ReportConfig, ReportValidationError


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / "data" / "exemplos" / "BANCO DE DADOS CV - AMANDA E EMILENE.xlsx"


def _write_csv(path: Path, content: str) -> Path:
    path.write_text(content, encoding="utf-8")
    return path


def _write_xlsx(path: Path, rows: list[list[object]]) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def test_reference_workbook_has_expected_tables_and_map_contract() -> None:
    result = ReportService().inspect(REFERENCE)

    assert len(result.sheets) == 21
    assert len(result.tables) == 22
    assert len(result.valid_tables) == 22
    assert sum(table.map_ready for table in result.tables) == 17
    assert any(table.source_sheet.strip() == "POLIOMIELITE" for table in result.tables)
    assert result.formula_count == 160


def test_csv_preserves_blank_and_zero(tmp_path: Path) -> None:
    path = _write_csv(
        tmp_path / "dados.csv",
        "Regiao;2020;2021;2022\nA;0;;2,5\nB;1;2;3\n",
    )
    result = ReportService().inspect(path)

    assert len(result.tables) == 1
    table = result.tables[0]
    assert table.valid
    assert table.years == (2020, 2021, 2022)
    assert table.series[0].values == (0.0, None, 2.5)


def test_invalid_value_is_exposed_and_blocks_report(tmp_path: Path) -> None:
    path = _write_csv(
        tmp_path / "invalido.csv",
        "Regiao;2020;2021;2022\nA;1;nao informado;3\nB;2;3;4\n",
    )
    service = ReportService()
    result = service.inspect(path)
    table = result.tables[0]

    assert not table.valid
    assert any(diagnostic.code == "INVALID_VALUE" for diagnostic in table.diagnostics)
    with pytest.raises(ReportValidationError):
        service.preview(result, ReportConfig((table.table_id,), years=(2020, 2021, 2022), title="Teste"))


def test_formula_without_cached_value_is_warning(tmp_path: Path) -> None:
    path = tmp_path / "formulas.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Indicador"
    sheet.append(["Regiao", 2020, 2021, 2022])
    sheet.append(["A", 1, "=1+1", 3])
    sheet.append(["B", 2, 3, 4])
    workbook.save(path)

    result = ReportService().inspect(path)
    table = result.tables[0]

    assert table.valid
    assert table.series[0].values == (1.0, None, 3.0)
    assert any(diagnostic.code == "FORMULA_NO_CACHED_VALUE" for diagnostic in table.diagnostics)


def test_multiple_tables_in_one_sheet_are_detected(tmp_path: Path) -> None:
    path = tmp_path / "duas_tabelas.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Dados"
    for row in (
        ["Primeira tabela", 2020, 2021, 2022],
        ["A", 1, 2, 3],
        ["B", 2, 3, 4],
        [],
        ["Segunda tabela", 2020, 2021, 2022],
        ["C", 4, 5, 6],
    ):
        sheet.append(row)
    workbook.save(path)

    result = ReportService().inspect(path)

    assert len(result.tables) == 2
    assert [table.header_row for table in result.tables] == [1, 5]


def test_remap_keeps_multiple_table_boundaries(tmp_path: Path) -> None:
    path = _write_xlsx(
        tmp_path / "duas_tabelas_remapeaveis.xlsx",
        [
            ["Primeira", 2020, 2021, 2022],
            ["A", 1, 2, 3],
            ["B", 4, 5, 6],
            [],
            ["Segunda", 2020, 2021, 2022],
            ["C", 7, 8, 9],
        ],
    )
    service = ReportService()
    result = service.inspect(path)
    first, second = result.tables

    remapped_first = service.remap(result, first.table_id, header_row=1, label_column=1)
    remapped_second = service.remap(remapped_first, second.table_id, header_row=5, label_column=1)

    assert [series.name for series in remapped_second.table(first.table_id).series] == ["A", "B"]
    assert [series.name for series in remapped_second.table(second.table_id).series] == ["C"]


def test_remap_rejects_year_column_as_series_identifier(tmp_path: Path) -> None:
    path = _write_csv(
        tmp_path / "coluna_ano.csv",
        "Regiao;2020;2021;2022\nA;1;2;3\nB;4;5;6\n",
    )
    service = ReportService()
    result = service.inspect(path)

    remapped = service.remap(result, result.tables[0].table_id, label_column=2)

    table = remapped.tables[0]
    assert not table.valid
    assert any(diagnostic.code == "LABEL_COLUMN_IS_YEAR" for diagnostic in table.diagnostics)


def test_data_values_that_look_like_years_do_not_create_header(tmp_path: Path) -> None:
    path = _write_csv(
        tmp_path / "valores_que_parecem_anos.csv",
        "Regiao;2019;2020;2021\nA;2020;2021;2022\n",
    )

    result = ReportService().inspect(path)

    assert len(result.tables) == 1
    assert result.tables[0].valid
    assert result.tables[0].series[0].values == (2020.0, 2021.0, 2022.0)


def test_validate_rejects_repeated_or_unsorted_years_and_tables(tmp_path: Path) -> None:
    path = _write_csv(
        tmp_path / "validacao.csv",
        "Regiao;2020;2021;2022\nA;1;2;3\n",
    )
    service = ReportService()
    result = service.inspect(path)
    table_id = result.tables[0].table_id

    diagnostics = service.validate(
        result,
        ReportConfig((table_id, table_id), years=(2021, 2020, 2020), title="Teste"),
    )
    codes = {diagnostic.code for diagnostic in diagnostics}

    assert {"DUPLICATE_TABLE_SELECTED", "YEARS_NOT_ORDERED"}.issubset(codes)


def test_methodology_uses_effective_years_unit_and_indicator(tmp_path: Path) -> None:
    path = _write_csv(
        tmp_path / "metodologia.csv",
        "Regiao;2020;2021;2022;2023\nA;1;2;3;4\n",
    )
    service = ReportService()
    result = service.inspect(path)
    table = result.tables[0]
    config = ReportConfig(
        (table.table_id,),
        "linhas",
        (2021, 2023),
        "Relatorio configurado",
        "",
        "",
        "Indicador configurado",
        "unidade escolhida",
    )

    artifact = service.preview(result, config)
    text = "\n".join(page.extract_text() or "" for page in PdfReader(str(artifact.pdf_path)).pages)
    service.cleanup(artifact)

    assert "Anos selecionados: 2021, 2023" in text
    assert "unidade escolhida" in text
    assert "Indicador configurado" in text


def test_csv_preserves_multiline_field(tmp_path: Path) -> None:
    path = _write_csv(
        tmp_path / "campo_multilinha.csv",
        'Regiao;2020;2021;2022\n"Regiao\nNorte";1;2;3\nSul;4;5;6\n',
    )

    result = ReportService().inspect(path)

    assert result.tables[0].series[0].name == "Regiao\nNorte"


def test_invalid_map_table_is_not_map_ready(tmp_path: Path) -> None:
    rows = ["Regiao;2020;2021;2022"]
    codes = ["4210", "4211", "4213", "4214", "4215", "4216", "4217", "4218"]
    rows.extend(f"{code} REGIAO;1;nao;3" for code in codes)
    path = _write_csv(tmp_path / "mapa_invalido.csv", "\n".join(rows) + "\n")

    table = ReportService().inspect(path).tables[0]

    assert not table.valid
    assert not table.map_ready


@pytest.mark.parametrize(
    "filename,content",
    [
        ("vacina_horizontal.csv", "Macrorregiao;2020;2021;2022\n4210 SUL;25;;0\n4211 NORTE;95;;105\n"),
        ("vacina_transposta.csv", "Ano;4210 SUL;4211 NORTE\n2020;25;95\n2021;;\n2022;0;105\n"),
        ("vacina_consolidada.csv", "Indicador;Macrorregiao;Ano;Cobertura\nVACINA ALFA;4210 SUL;2020;25\nVACINA ALFA;4211 NORTE;2020;95\nVACINA ALFA;4210 SUL;2021;\nVACINA ALFA;4211 NORTE;2021;\nVACINA ALFA;4210 SUL;2022;0\nVACINA ALFA;4211 NORTE;2022;105\n"),
    ],
)
def test_partial_map_marks_absent_regions_and_rejects_empty_year(tmp_path: Path, filename: str, content: str) -> None:
    path = _write_csv(tmp_path / filename, content)
    service = ReportService()
    result = service.inspect(path)
    assert len(result.valid_tables) == 1
    table = result.valid_tables[0]
    assert table.map_ready
    assert table.to_dict()["regioes_mapeadas"] == 2
    assert table.series[0].values[table.years.index(2022)] == 0

    empty_year = ReportConfig((table.table_id,), "mapa", (2021,), "Mapa parcial")
    assert any(item.code == "MAP_YEAR_WITHOUT_DATA" for item in service.validate(result, empty_year))

    valid = ReportConfig((table.table_id,), "mapa", (2022,), "Mapa parcial")
    artifact = service.preview(result, valid)
    try:
        page = PdfReader(str(artifact.pdf_path)).pages[0].extract_text()
        assert artifact.page_count == 2
        assert "Sem dado" in page
        assert "0,00%" in page
        assert "105,00%" in page
        assert "nenhum valor foi estimado" in page
    finally:
        service.cleanup(artifact)


def test_abandonment_rate_with_sc_codes_does_not_use_coverage_map(tmp_path: Path) -> None:
    path = _write_csv(tmp_path / "taxa_de_abandono_da_vacina.csv", "Macrorregiao;2020;2021;2022\n4210 SUL;12;10;9\n4211 NORTE;20;18;17\n")
    table = ReportService().inspect(path).valid_tables[0]
    assert table.indicator_type == "desconhecido"
    assert table.unit == "Taxa informada na planilha"
    assert not table.map_ready


def test_export_cancellation_does_not_publish_or_leave_temporary_file(tmp_path: Path) -> None:
    source = _write_csv(
        tmp_path / "dados.csv",
        "Regiao;2020;2021;2022\nA;1;2;3\n",
    )
    service = ReportService()
    result = service.inspect(source)
    table = result.tables[0]
    artifact = service.preview(result, ReportConfig((table.table_id,), years=(2020, 2021, 2022), title="Teste"))
    cancel_event = threading.Event()

    def cancel_after_copy_starts(value: int, message: str) -> None:
        if message == "Copiando o PDF para o destino":
            cancel_event.set()

    destination = tmp_path / "saida" / "cancelado.pdf"
    with pytest.raises(OperationCancelled):
        service.export(artifact, destination, progress=cancel_after_copy_starts, cancel_event=cancel_event)

    assert not destination.exists()
    assert not list(destination.parent.glob(".*.tmp")) if destination.parent.exists() else True
    service.cleanup(artifact)


def test_preview_export_is_atomic_and_keeps_input_unchanged(tmp_path: Path) -> None:
    source = _write_csv(
        tmp_path / "dados com acento.csv",
        "Regiao;2020;2021;2022\nA;0;;2\nB;1;2;3\n",
    )
    service = ReportService()
    result = service.inspect(source)
    table = result.tables[0]
    before = source.read_bytes()
    config = ReportConfig((table.table_id,), "linhas", (2020, 2021, 2022), "Relatorio de teste", "Autoria", "Fonte", table.title, table.unit)
    artifact = service.preview(result, config)
    destination = tmp_path / "pasta com espaco" / "relatorio final.pdf"

    exported = service.export(artifact, destination)
    assert exported.exists()
    assert len(PdfReader(str(exported)).pages) == 2
    assert source.read_bytes() == before
    assert not list(destination.parent.glob(".*.tmp"))
    with pytest.raises(FileExistsError):
        service.export(artifact, destination)
    service.cleanup(artifact)


def test_map_preview_contains_methodology(tmp_path: Path) -> None:
    del tmp_path
    service = ReportService()
    result = service.inspect(REFERENCE)
    table = next(table for table in result.tables if table.source_sheet.strip() == "TRÍPLICE VIRAL")
    config = ReportConfig((table.table_id,), "mapa", (2025,), "Mapa de teste", "", "", table.title, table.unit)

    artifact = service.preview(result, config)
    text = "\n".join(page.extract_text() or "" for page in PdfReader(str(artifact.pdf_path)).pages)

    assert artifact.page_count == 2
    assert "Metodologia e origem dos dados" in text
    assert "limites municipais do IBGE" in text
    service.cleanup(artifact)
