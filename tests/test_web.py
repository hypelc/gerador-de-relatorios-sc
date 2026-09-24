from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
import openpyxl
from fastapi.testclient import TestClient
from pypdf import PdfReader

from gerador_sc.regional import read_values
from gerador_sc.importers import inspect_file
from webapp import api as api_module
from webapp.api import create_app

ROOT = Path(__file__).resolve().parents[1]
ANNUAL = ROOT / "data" / "exemplos" / "BANCO DE DADOS CV - AMANDA E EMILENE.xlsx"
REGIONAL = ROOT / "data" / "exemplos" / "taxa estado SC.xlsx"


@pytest.fixture(autouse=True)
def reset_test_rate_limits() -> None:
    api_module._rate_events.clear()


def client() -> TestClient:
    return TestClient(create_app())


def test_public_api_accepts_upload_without_session() -> None:
    web = client()
    response = web.post("/api/inspect", files={"file": ("dados.csv", b"a;b\n1;2")})
    assert response.status_code == 422
    assert response.json()["detail"]["message"] == "Nenhuma tabela válida foi encontrada."
    assert web.get("/api/session").status_code == 404


@pytest.mark.parametrize("variation,annual_count", [(1, 3), (2, 3), (3, 1), (4, 3), (5, 1)])
def test_five_fictional_layouts_generate_annual_and_regional_bar_pdfs(variation: int, annual_count: int) -> None:
    path = ROOT / "tests" / "fixtures" / f"dados_ficticios_variacao_{variation:02d}.xlsx"
    original = path.read_bytes()
    files = {"file": (path.name, original)}
    web = client()
    inspection = web.post("/api/inspect", files=files)
    assert inspection.status_code == 200, inspection.text
    data = inspection.json()
    annual = [table for table in data["import"]["tabelas"] if table["valida"]]
    categorical = data["import"]["tabelas_categoricas"]
    assert len(annual) == annual_count
    assert len(categorical) == 1
    assert all(table["anos"] == list(range(2016, 2026)) for table in annual)
    assert all(not table["mapa_disponivel"] for table in annual)
    assert len(categorical[0]["categorias"]) == 8

    options = {"kind": "annual", "title": "Comparação das regiões", "authors": "Equipe", "source": "Planilha fictícia", "table_ids": [annual[0]["id"]], "years": [2025], "model": "barras", "confirmed_exclusions": []}
    pdf = web.post("/api/generate", files=files, data={"options": json.dumps(options)})
    assert pdf.status_code == 200, pdf.text
    assert len(PdfReader(io.BytesIO(pdf.content)).pages) == 2
    options = {"kind": "annual", "title": "Taxa por regional", "authors": "Equipe", "source": "Planilha fictícia", "unit": "taxa", "period": "2025", "model": "barras_categoria", "category_id": categorical[0]["id"]}
    pdf = web.post("/api/generate", files=files, data={"options": json.dumps(options)})
    assert pdf.status_code == 200, pdf.text
    pages = PdfReader(io.BytesIO(pdf.content)).pages
    assert len(pages) == 2
    assert "2025" in pages[-1].extract_text()
    assert path.read_bytes() == original


def test_pie_rejects_rates_and_preserves_zero_and_over_100() -> None:
    path = ROOT / "tests" / "fixtures" / "dados_ficticios_variacao_01.xlsx"
    web = client()
    files = {"file": (path.name, path.read_bytes())}
    table = web.post("/api/inspect", files=files).json()["import"]["tabelas"][0]
    options = {"kind": "annual", "title": "Tentativa", "table_ids": [table["id"]], "years": [2025], "model": "pizza", "confirmed_exclusions": []}
    response = web.post("/api/generate", files=files, data={"options": json.dumps(options)})
    assert response.status_code == 422
    assert any(item["code"] == "PIE_REQUIRES_COUNTS" for item in response.json()["detail"]["diagnostics"])
    assert any(value is not None and value > 100 for series in table["series"] for value in series["valores"])


def test_pie_generates_for_regional_birth_counts() -> None:
    web = client()
    files = {"file": (ANNUAL.name, ANNUAL.read_bytes())}
    tables = web.post("/api/inspect", files=files).json()["import"]["tabelas"]
    table = next(item for item in tables if item["aba"] == "NASCIDOS VIVOS SC")
    options = {"kind": "annual", "title": "Nascidos vivos por região", "table_ids": [table["id"]], "years": [2025], "model": "pizza", "confirmed_exclusions": []}
    response = web.post("/api/generate", files=files, data={"options": json.dumps(options)})
    assert response.status_code == 200, response.text
    assert len(PdfReader(io.BytesIO(response.content)).pages) == 2


def test_scrambled_years_keep_original_cell_values_and_unknown_unit() -> None:
    path = ROOT / "tests" / "fixtures" / "dados_ficticios_variacao_05.xlsx"
    result = inspect_file(path)
    table = result.valid_tables[0]
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        first_original = workbook.active.cell(6, 4).value
        label_original = workbook.active.cell(6, 5).value
    finally:
        workbook.close()
    assert table.series[0].name == label_original
    assert table.series[0].values[table.years.index(2025)] == first_original
    assert table.indicator_type == "desconhecido"
    assert table.unit == "Valor informado na planilha"


def test_regional_workbook_generates_pdf_and_png_without_changing_source() -> None:
    web = client()
    source = REGIONAL.read_bytes()
    files = {"file": (REGIONAL.name, source)}
    inspected = web.post("/api/inspect", files=files)
    assert inspected.status_code == 200
    data = inspected.json()
    assert data["kind"] == "regional"
    assert len(data["regions"]) == 17
    options = {
        "kind": "regional",
        "title": "Distribuição da taxa de profissionais",
        "authors": "Equipe de pesquisa",
        "source": "Planilha fornecida",
        "period": "2025",
        "unit": "por 1.000 habitantes",
        "format": "pdf",
    }
    pdf = web.post("/api/generate", files=files, data={"options": json.dumps(options)})
    assert pdf.status_code == 200
    assert pdf.headers["content-type"].startswith("application/pdf")
    text = "\n".join(
        page.extract_text() for page in PdfReader(io.BytesIO(pdf.content)).pages
    )
    assert "Equipe de pesquisa" in text
    assert "2025" in text
    assert "sem estimativa municipal" in text
    options["format"] = "png"
    png = web.post("/api/generate", files=files, data={"options": json.dumps(options)})
    assert png.status_code == 200
    assert png.content.startswith(b"\x89PNG")
    assert REGIONAL.read_bytes() == source


def test_annual_workbook_generates_web_pdf() -> None:
    web = client()
    files = {"file": (ANNUAL.name, ANNUAL.read_bytes())}
    inspected = web.post("/api/inspect", files=files)
    assert inspected.status_code == 200
    data = inspected.json()
    assert data["kind"] == "annual"
    assert len(data["import"]["tabelas"]) == 22
    table = next(table for table in data["import"]["tabelas"] if table["valida"])
    options = {
        "kind": "annual",
        "title": "Evolução da cobertura",
        "authors": "Equipe",
        "source": "Planilha da pesquisa",
        "unit": "",
        "table_ids": [table["id"]],
        "years": table["anos"],
        "model": "paineis",
        "confirmed_exclusions": [],
    }
    pdf = web.post("/api/generate", files=files, data={"options": json.dumps(options)})
    assert pdf.status_code == 200
    pages = PdfReader(io.BytesIO(pdf.content)).pages
    assert len(pages) == 2
    method = pages[-1].extract_text()
    assert "Processamento no servidor" in method
    assert "Processamento local" not in method


def test_mismatched_model_is_rejected() -> None:
    web = client()
    options = {
        "kind": "annual",
        "title": "Erro",
        "table_ids": [],
        "years": [],
        "model": "paineis",
    }
    response = web.post(
        "/api/generate",
        files={"file": (REGIONAL.name, REGIONAL.read_bytes())},
        data={"options": json.dumps(options)},
    )
    assert response.status_code == 422


def test_regional_csv_is_recognized_and_missing_region_is_explained() -> None:
    web = client()
    regions, state, _, _ = read_values(REGIONAL)
    rows = [
        "Regionais;Taxa de profissionais",
        *(f"{name};{value}" for name, value in regions.values()),
        f"Santa Catarina;{state}",
    ]
    csv_data = "\n".join(rows).encode("utf-8")
    recognized = web.post("/api/inspect", files={"file": ("regionais.csv", csv_data)})
    assert recognized.status_code == 200
    assert recognized.json()["kind"] == "regional"
    assert len(recognized.json()["regions"]) == 17
    missing = web.post(
        "/api/inspect",
        files={"file": ("regionais.csv", "\n".join(rows[:2] + rows[-1:]).encode())},
    )
    assert missing.status_code == 422
    assert "Faltam Regionais" in missing.json()["detail"]


def test_invalid_spreadsheets_have_readable_errors() -> None:
    web = client()
    malformed_csv = web.post(
        "/api/inspect", files={"file": ("dados.csv", b"Regionais;Taxa\n\xff;1")}
    )
    assert malformed_csv.status_code == 422
    assert "CSV" in malformed_csv.json()["detail"]

    fake_xlsx = web.post(
        "/api/inspect", files={"file": ("dados.xlsx", b"not an Excel workbook")}
    )
    assert fake_xlsx.status_code == 400
    assert "danificado" in fake_xlsx.json()["detail"]
